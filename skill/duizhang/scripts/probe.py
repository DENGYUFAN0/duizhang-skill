#!/usr/bin/env python3
"""
probe.py —— 生成探针文件

对若干个待验证的候选取值做笛卡尔积，每种组合分配样本数据里一行"本来就要导入"
的真实数据，把组合值覆盖进对应列后写进模板副本；同时输出一张
"输出行号 -> 各列取值 -> 样本行号" 的对照表，方便导入平台返回结果后逐行对号。

注意：探针文件里能通过校验的行，会真的写入生产系统，所以只应使用本来就要
导入的真实数据去做排列组合，不要编造数据。

用法示例：
    python probe.py --template empty.xlsx --sheet Sheet1 --header-rows 4 \
        --sample real_rows.xlsx --vary vary.json --pk A --out probe.xlsx --map-out map.csv

vary.json 形如：
    {"C": ["候选值1", "候选值2"], "E": ["候选值A", "候选值B", "候选值C"]}
"""

import argparse
import itertools
import json
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

sys.stdout.reconfigure(encoding="utf-8")


def parse_col_spec(spec: str) -> int:
    spec = spec.strip()
    if not spec:
        raise ValueError("列标识不能为空")
    if spec.isdigit():
        n = int(spec)
        if n < 1:
            raise ValueError(f"列序号必须从 1 开始，收到：{spec}")
        return n - 1
    return column_index_from_string(spec.upper()) - 1


def read_sample(path: str, sheet, header_rows: int) -> pd.DataFrame:
    lower = path.lower()
    if lower.endswith(".csv"):
        df = pd.read_csv(path, header=None, skiprows=header_rows, dtype=str)
    else:
        sheet_name = sheet if sheet is not None else 0
        df = pd.read_excel(
            path, sheet_name=sheet_name, header=None, skiprows=header_rows,
            dtype=str, engine="openpyxl",
        )
    df = df.dropna(how="all").reset_index(drop=True)
    return df


def main():
    parser = argparse.ArgumentParser(
        description="生成探针文件：用真实样本行承载候选值的笛卡尔积组合，写入模板副本供黑盒平台导入"
    )
    parser.add_argument("--template", required=True, help="空模板 xlsx（保留其表头行与格式）")
    parser.add_argument("--sheet", required=True, help="工作表名")
    parser.add_argument("--header-rows", type=int, required=True, help="表头行数（模板与样本文件通用）")
    parser.add_argument("--sample", required=True, help="样本数据文件（.csv 或 .xlsx），含本来就要导入的真实行")
    parser.add_argument("--vary", required=True, help="JSON 文件：{\"列字母\": [候选值1, 候选值2, ...], ...}")
    parser.add_argument("--pk", required=True, help="主键列（字母或序号）")
    parser.add_argument("--out", required=True, help="输出的探针 xlsx 路径")
    parser.add_argument("--map-out", required=True, help="输出的对照表路径（.csv 或 .xlsx）")
    args = parser.parse_args()

    print("提醒：探针里能通过校验的行会真的写进生产系统，所以只用本来就要导入的真实数据做组合，不要编造数据。")

    with open(args.vary, "r", encoding="utf-8") as f:
        vary_raw = json.load(f)
    if not vary_raw:
        print("错误：--vary 指定的 JSON 为空", file=sys.stderr)
        sys.exit(1)

    try:
        vary_cols = {parse_col_spec(k): v for k, v in vary_raw.items()}
        pk_idx = parse_col_spec(args.pk)
    except ValueError as e:
        print(f"错误：列标识解析失败：{e}", file=sys.stderr)
        sys.exit(1)

    col_letters = list(vary_raw.keys())
    col_idx_ordered = list(vary_cols.keys())
    value_lists = [vary_cols[i] for i in col_idx_ordered]

    combos = list(itertools.product(*value_lists))
    n_combos = len(combos)
    print(f"候选值笛卡尔积组合数：{n_combos}")

    sample_df = read_sample(args.sample, args.sheet, args.header_rows)
    n_sample = len(sample_df)
    print(f"样本可用真实行数：{n_sample}")

    if n_combos > n_sample:
        print(
            f"错误：组合数（{n_combos}）超过样本行数（{n_sample}），"
            f"没有足够多本来就要导入的真实行来承载每个组合，无法继续。",
            file=sys.stderr,
        )
        sys.exit(1)

    max_needed_col = max([pk_idx] + col_idx_ordered)
    if max_needed_col >= sample_df.shape[1]:
        print(f"错误：样本数据只有 {sample_df.shape[1]} 列，但涉及的列超出范围", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(args.template)
    if args.sheet not in wb.sheetnames:
        print(f"错误：模板中没有名为 {args.sheet} 的工作表", file=sys.stderr)
        sys.exit(1)
    ws = wb[args.sheet]

    start_row = args.header_rows + 1  # Excel 1-based 行号，紧跟表头之后
    n_cols = sample_df.shape[1]

    map_rows = []
    used_pk_values = []

    for i, combo in enumerate(combos):
        sample_row = sample_df.iloc[i]
        out_row_num = start_row + i

        # 先把整行真实数据原样写入
        for c in range(n_cols):
            ws.cell(row=out_row_num, column=c + 1, value=sample_row.iloc[c])

        # 再把这个组合要验证的取值覆盖进对应列
        combo_values = {}
        for col_i, val in zip(col_idx_ordered, combo):
            ws.cell(row=out_row_num, column=col_i + 1, value=val)
            combo_values[col_i] = val

        pk_value = combo_values.get(pk_idx, sample_row.iloc[pk_idx])
        used_pk_values.append(pk_value)

        map_row = {"输出行号": out_row_num}
        for letter, col_i in zip(col_letters, col_idx_ordered):
            map_row[f"列{letter}取值"] = combo_values[col_i]
        map_row["样本行号"] = i + 1  # 样本数据（去表头、去空行后）从 1 开始计的行号
        map_rows.append(map_row)

    # 校验主键不重复
    if len(set(used_pk_values)) != len(used_pk_values):
        print("错误：生成的探针行里主键出现重复值，无法保证每行都能被平台区分开，已中止写出。", file=sys.stderr)
        sys.exit(1)

    wb.save(args.out)
    print(f"已写出探针文件：{args.out}（共 {n_combos} 行数据，从第 {start_row} 行开始）")

    map_df = pd.DataFrame(map_rows)
    if args.map_out.lower().endswith(".xlsx"):
        map_df.to_excel(args.map_out, index=False)
    else:
        map_df.to_csv(args.map_out, index=False, encoding="utf-8-sig")
    print(f"已写出对照表：{args.map_out}")


if __name__ == "__main__":
    main()
