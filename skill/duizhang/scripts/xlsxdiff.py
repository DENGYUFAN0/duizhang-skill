#!/usr/bin/env python3
"""
xlsxdiff.py —— 两个 xlsx 的结构对照

拿一个"导入成功"的文件和一个"导入失败"的文件做结构级对照（不是内容级的逐行
diff），把两边在 sheet 结构、表头、合并单元格、数据验证、公式、各列统计特征
上的差异列出来，帮助定位黑盒平台到底在挑哪个结构性毛病。只输出不同之处。

用法示例：
    python xlsxdiff.py good.xlsx bad.xlsx --header-rows 4
"""

import argparse
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

ZERO_WIDTH_CHARS = ["​", "‌", "‍", "﻿", "⁠"]
MAIN_NS = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def get_sheet_order_from_xml(path: str):
    """直接读 xl/workbook.xml 里 <sheets> 的顺序，不经 openpyxl 的二次整理。"""
    try:
        with zipfile.ZipFile(path) as z:
            with z.open("xl/workbook.xml") as f:
                tree = ET.parse(f)
        root = tree.getroot()
        sheets_el = root.find("main:sheets", MAIN_NS)
        if sheets_el is None:
            return None
        return [s.attrib.get("name") for s in sheets_el.findall("main:sheet", MAIN_NS)]
    except Exception:
        return None


def has_whitespace_issue(v) -> bool:
    if not isinstance(v, str):
        return False
    if v != v.strip():
        return True
    if "\n" in v or "\r" in v:
        return True
    if any(ch in v for ch in ZERO_WIDTH_CHARS):
        return True
    return False


def classify_type(v) -> str:
    if v is None:
        return "None"
    if isinstance(v, (datetime, date, time)):
        return "datetime"
    if isinstance(v, bool):
        return "int"  # Python 的 bool 是 int 子类，归入 int
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "float"
    if isinstance(v, str):
        return "str"
    return "其他"


def get_column_values(ws, col_idx: int):
    """取某个 1-based 列在整个 sheet（含表头行）里的全部单元格值。"""
    if ws is None or col_idx > ws.max_column:
        return []
    return [ws.cell(row=r, column=col_idx).value for r in range(1, ws.max_row + 1)]


def count_formula_cells(ws) -> int:
    if ws is None:
        return 0
    n = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                n += 1
    return n


def count_whitespace_cells(ws) -> int:
    if ws is None:
        return 0
    n = 0
    for row in ws.iter_rows():
        for cell in row:
            if has_whitespace_issue(cell.value):
                n += 1
    return n


def get_merged_ranges(ws):
    if ws is None:
        return set()
    return {str(r) for r in ws.merged_cells.ranges}


def get_data_validations(ws):
    if ws is None:
        return set()
    result = set()
    for dv in ws.data_validations.dataValidation:
        result.add((str(dv.sqref), str(dv.type), str(dv.operator), str(dv.formula1), str(dv.formula2)))
    return result


def main():
    parser = argparse.ArgumentParser(
        description="两个 xlsx 的结构对照：只列出结构性差异，用于定位导入成功/失败文件之间的结构性区别"
    )
    parser.add_argument("file_a", help="文件 A（建议：导入成功的文件）")
    parser.add_argument("file_b", help="文件 B（建议：导入失败的文件）")
    parser.add_argument("--header-rows", type=int, default=1, help="表头行数，用于逐格比对表头（默认 1）")
    args = parser.parse_args()

    print(f"文件 A：{args.file_a}")
    print(f"文件 B：{args.file_b}")
    print(f"表头行数：{args.header_rows}")
    print()

    any_diff = False

    def report(line: str):
        nonlocal any_diff
        any_diff = True
        print(line)

    sheets_a = get_sheet_order_from_xml(args.file_a)
    sheets_b = get_sheet_order_from_xml(args.file_b)

    wb_a = load_workbook(args.file_a, data_only=False)
    wb_b = load_workbook(args.file_b, data_only=False)

    if sheets_a is None:
        sheets_a = list(wb_a.sheetnames)
    if sheets_b is None:
        sheets_b = list(wb_b.sheetnames)

    print("=== sheet 名与顺序 ===")
    if sheets_a != sheets_b:
        report(f"A: {sheets_a}")
        report(f"B: {sheets_b}")
        set_a, set_b = set(sheets_a), set(sheets_b)
        if set_a - set_b:
            report(f"只在 A 出现的 sheet：{sorted(set_a - set_b)}")
        if set_b - set_a:
            report(f"只在 B 出现的 sheet：{sorted(set_b - set_a)}")
        if set_a == set_b:
            report("两边 sheet 集合相同，但顺序不同")
    else:
        print("（一致）")
    print()

    common_sheets = [s for s in sheets_a if s in sheets_b]

    for sheet in common_sheets:
        ws_a = wb_a[sheet]
        ws_b = wb_b[sheet]
        sheet_header_printed = False

        def sheet_report(line: str):
            nonlocal sheet_header_printed
            if not sheet_header_printed:
                print(f"--- sheet: {sheet} ---")
                sheet_header_printed = True
            report(line)

        # dimension
        if ws_a.dimensions != ws_b.dimensions:
            sheet_report(f"[dimension] A={ws_a.dimensions}  B={ws_b.dimensions}")

        # 表头逐格
        max_col_header = max(ws_a.max_column, ws_b.max_column)
        for r in range(1, args.header_rows + 1):
            for c in range(1, max_col_header + 1):
                va = ws_a.cell(row=r, column=c).value if r <= ws_a.max_row and c <= ws_a.max_column else None
                vb = ws_b.cell(row=r, column=c).value if r <= ws_b.max_row and c <= ws_b.max_column else None
                if va != vb:
                    coord = f"{get_column_letter(c)}{r}"
                    sheet_report(f"[表头 {coord}] A={va!r}  B={vb!r}")

        # 合并单元格
        merged_a = get_merged_ranges(ws_a)
        merged_b = get_merged_ranges(ws_b)
        if merged_a != merged_b:
            only_a = merged_a - merged_b
            only_b = merged_b - merged_a
            if only_a:
                sheet_report(f"[合并单元格] 只在 A：{sorted(only_a)}")
            if only_b:
                sheet_report(f"[合并单元格] 只在 B：{sorted(only_b)}")

        # 数据验证
        dv_a = get_data_validations(ws_a)
        dv_b = get_data_validations(ws_b)
        if dv_a != dv_b:
            only_a = dv_a - dv_b
            only_b = dv_b - dv_a
            if only_a:
                sheet_report(f"[数据验证] 只在 A：{sorted(only_a)}")
            if only_b:
                sheet_report(f"[数据验证] 只在 B：{sorted(only_b)}")

        # 公式格数
        f_a = count_formula_cells(ws_a)
        f_b = count_formula_cells(ws_b)
        if f_a != f_b:
            sheet_report(f"[公式格数] A={f_a}  B={f_b}")

        # 首尾空白/换行/零宽字符格数（整个 sheet 汇总）
        ws_count_a = count_whitespace_cells(ws_a)
        ws_count_b = count_whitespace_cells(ws_b)
        if ws_count_a != ws_count_b:
            sheet_report(f"[首尾空白/换行/零宽字符格数] A={ws_count_a}  B={ws_count_b}")

        # 每列统计
        max_col = max(ws_a.max_column, ws_b.max_column)
        for c in range(1, max_col + 1):
            letter = get_column_letter(c)
            values_a = get_column_values(ws_a, c)
            values_b = get_column_values(ws_b, c)

            non_empty_a = sum(1 for v in values_a if v is not None and v != "")
            non_empty_b = sum(1 for v in values_b if v is not None and v != "")
            if non_empty_a != non_empty_b:
                sheet_report(f"[列{letter} 非空数] A={non_empty_a}  B={non_empty_b}")

            len_a = max((len(str(v)) for v in values_a if v is not None), default=0)
            len_b = max((len(str(v)) for v in values_b if v is not None), default=0)
            if len_a != len_b:
                sheet_report(f"[列{letter} 最大文本长度] A={len_a}  B={len_b}")

            dist_a = dict(Counter(classify_type(v) for v in values_a))
            dist_b = dict(Counter(classify_type(v) for v in values_b))
            if dist_a != dist_b:
                sheet_report(f"[列{letter} 值类型分布] A={dist_a}  B={dist_b}")

        if sheet_header_printed:
            print()

    if not any_diff:
        print("在所检查的结构维度上，两个文件没有发现差异。")


if __name__ == "__main__":
    main()
