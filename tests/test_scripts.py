#!/usr/bin/env python3
"""
test_scripts.py —— duizhang-skill 四个脚本的验证测试

优先用 pytest 跑（自动发现 test_ 开头的函数）；没装 pytest 时，
直接 `python tests\\test_scripts.py` 也能跑（文件末尾的 __main__ 会
手动依次调用每个 test_ 函数并汇总结果）。

所有测试都在临时目录里用 openpyxl 现场生成合成数据，不依赖任何外部文件。
"""

import csv
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

sys.stdout.reconfigure(encoding="utf-8")

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "skill" / "duizhang" / "scripts"


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def run_script(script_name: str, args):
    cmd = [sys.executable, str(SCRIPTS_DIR / script_name)] + [str(a) for a in args]
    return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")


# ---------------------------------------------------------------------------
# dupcount.py
# ---------------------------------------------------------------------------

def _make_case_insensitive_dataset(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name"])  # 表头，1 行
    rows = [
        (1, "abc"),
        (2, "ABC"),
        (3, "xyz"),
        (4, "XYZ"),
        (5, "hello"),
    ]
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_dupcount_case_insensitive_vs_exact():
    dupcount = load_module(SCRIPTS_DIR / "dupcount.py", "dupcount_mod_1")
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "data.xlsx"
        _make_case_insensitive_dataset(path)

        df = dupcount.read_table(str(path), "Sheet1", 1)
        assert len(df) == 5, "跳过表头后应有 5 行数据"

        key_idx = [1]  # Name 列，0-based 索引 1

        composite_exact = dupcount.build_composite(df, key_idx, dupcount.cipher_exact)
        assert int(composite_exact.duplicated(keep="first").sum()) == 0, "精确口径不应有重复"

        composite_ci = dupcount.build_composite(df, key_idx, dupcount.cipher_lower)
        assert int(composite_ci.duplicated(keep="first").sum()) == 2, "忽略大小写口径 duplicated() 应为 2"
        assert int(composite_ci.duplicated(keep=False).sum()) == 4, "忽略大小写口径 duplicated(keep=False) 应为 4"


def test_dupcount_target_hit_via_cli():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "data.xlsx"
        _make_case_insensitive_dataset(path)

        result = run_script(
            "dupcount.py",
            [path, "--sheet", "Sheet1", "--header-rows", "1", "--keys", "B", "--target", "2"],
        )
        assert result.returncode == 0, f"dupcount.py 运行失败: {result.stderr}"
        assert "忽略大小写" in result.stdout
        assert "命中" in result.stdout


# ---------------------------------------------------------------------------
# combos.py
# ---------------------------------------------------------------------------

def test_combos_diff_sets():
    combos_mod = load_module(SCRIPTS_DIR / "combos.py", "combos_mod_1")
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        fa = td / "a.xlsx"
        fb = td / "b.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["K1", "K2"])
        ws.append(["x", "1"])
        ws.append(["y", "2"])
        ws.append(["common", "9"])
        wb.save(fa)

        wb2 = Workbook()
        ws2 = wb2.active
        ws2.title = "Sheet1"
        ws2.append(["K1", "K2"])
        ws2.append(["z", "3"])
        ws2.append(["common", "9"])
        wb2.save(fb)

        df_a = combos_mod.read_table(str(fa), "Sheet1", 1)
        df_b = combos_mod.read_table(str(fb), "Sheet1", 1)
        col_idx = combos_mod.parse_col_list("A,B")
        counts_a = combos_mod.combo_counts(df_a, col_idx)
        counts_b = combos_mod.combo_counts(df_b, col_idx)

        only_a = set(counts_a) - set(counts_b)
        only_b = set(counts_b) - set(counts_a)

        assert only_a == {("x", "1"), ("y", "2")}
        assert only_b == {("z", "3")}
        assert ("common", "9") not in only_a
        assert ("common", "9") not in only_b


# ---------------------------------------------------------------------------
# probe.py
# ---------------------------------------------------------------------------

def _make_probe_inputs(td: Path, n_sample_rows: int):
    template_path = td / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["PK", "Name", "Status"])  # 表头 1 行
    wb.save(template_path)

    sample_path = td / "sample.csv"
    with open(sample_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["PK", "Name", "Status"])  # 表头 1 行，与 header_rows 一致
        for i in range(1, n_sample_rows + 1):
            w.writerow([i, f"name{i}", "orig"])

    vary_path = td / "vary.json"
    vary_path.write_text(
        json.dumps({"B": ["v1", "v2"], "C": ["s1", "s2"]}, ensure_ascii=False),
        encoding="utf-8",
    )
    return template_path, sample_path, vary_path


def test_probe_generates_correct_rows_and_unique_pk():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        template_path, sample_path, vary_path = _make_probe_inputs(td, n_sample_rows=5)
        out_path = td / "probe_out.xlsx"
        map_path = td / "map.csv"

        result = run_script(
            "probe.py",
            [
                "--template", template_path, "--sheet", "Sheet1", "--header-rows", 1,
                "--sample", sample_path, "--vary", vary_path, "--pk", "A",
                "--out", out_path, "--map-out", map_path,
            ],
        )
        assert result.returncode == 0, f"probe.py 运行失败: {result.stderr}"

        wb_out = load_workbook(out_path)
        ws_out = wb_out["Sheet1"]
        data_rows = [
            r for r in ws_out.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in r)
        ]
        assert len(data_rows) == 4, "候选值笛卡尔积组合数应为 2*2=4，对应 4 行数据"

        pk_values = [r[0] for r in data_rows]
        assert len(set(pk_values)) == len(pk_values), "主键不应重复"

        map_df = pd.read_csv(map_path, encoding="utf-8-sig")
        assert len(map_df) == 4, "对照表行数应与组合数一致"


def test_probe_errors_when_sample_rows_insufficient():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        # 组合数是 2*2=4，样本只给 3 行，应该报错退出
        template_path, sample_path, vary_path = _make_probe_inputs(td, n_sample_rows=3)
        out_path = td / "probe_out.xlsx"
        map_path = td / "map.csv"

        result = run_script(
            "probe.py",
            [
                "--template", template_path, "--sheet", "Sheet1", "--header-rows", 1,
                "--sample", sample_path, "--vary", vary_path, "--pk", "A",
                "--out", out_path, "--map-out", map_path,
            ],
        )
        assert result.returncode != 0, "样本行数不够时应该报错退出，而不是静默生成"
        assert not out_path.exists(), "样本不够时不应该写出探针文件"


# ---------------------------------------------------------------------------
# xlsxdiff.py
# ---------------------------------------------------------------------------

def test_xlsxdiff_finds_formula_and_header_diff():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        fa = td / "a.xlsx"
        fb = td / "b.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["ID", "Amount"])
        ws.append([1, 10])
        ws.append([2, 20])
        wb.save(fa)

        wb2 = Workbook()
        ws2 = wb2.active
        ws2.title = "Sheet1"
        ws2.append(["ID ", "Amount"])  # 表头多了个尾随空格
        ws2.append([1, 10])
        ws2.append([2, "=1+1"])  # 公式格
        wb2.save(fb)

        result = run_script("xlsxdiff.py", [fa, fb, "--header-rows", 1])
        assert result.returncode == 0, f"xlsxdiff.py 运行失败: {result.stderr}"
        out = result.stdout
        assert "表头" in out, "应该发现表头差异"
        assert "公式格数" in out, "应该发现公式格数差异"


def test_xlsxdiff_no_diff_reports_clean():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        fa = td / "a.xlsx"
        fb = td / "b.xlsx"

        for p in (fa, fb):
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["ID", "Amount"])
            ws.append([1, 10])
            ws.append([2, 20])
            wb.save(p)

        result = run_script("xlsxdiff.py", [fa, fb, "--header-rows", 1])
        assert result.returncode == 0
        assert "没有发现差异" in result.stdout


# ---------------------------------------------------------------------------
# 支持直接 `python tests\test_scripts.py` 运行（不依赖 pytest）
# ---------------------------------------------------------------------------

_TEST_FUNCS = [
    obj for name, obj in list(globals().items())
    if name.startswith("test_") and callable(obj)
]

if __name__ == "__main__":
    failures = []
    for fn in _TEST_FUNCS:
        try:
            fn()
            print(f"[PASS] {fn.__name__}")
        except AssertionError as e:
            failures.append(fn.__name__)
            print(f"[FAIL] {fn.__name__}: {e}")
        except Exception as e:
            failures.append(fn.__name__)
            print(f"[ERROR] {fn.__name__}: {e}")

    print()
    if failures:
        print(f"共 {len(failures)} 个测试失败：{failures}")
        sys.exit(1)
    else:
        print(f"全部 {len(_TEST_FUNCS)} 个测试通过")
